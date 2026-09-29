"""Create the intentionally minimal DOCX source for conversion to genuine DOC."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
out = root / 'public' / 'resources'
out.mkdir(parents=True, exist_ok=True)
document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
normal = document.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
title = document.styles['Title']
title.font.name = 'Calibri'
title.font.size = Pt(24)
title.font.color.rgb = RGBColor(0, 0, 0)
document.add_paragraph('Didarian Publicating', style='Title')
document.add_paragraph('Template design in processing')
# The bundled base DOCX may include title borders. This deliberately plain
# resource uses whitespace only, so remove inherited/direct paragraph borders.
for element in (document._element, document.styles.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
document.core_properties.author = 'Didarian Publicating'
document.core_properties.last_modified_by = 'Didarian Publicating'
document.core_properties.title = 'Didarian Publicating'
document.core_properties.subject = 'Research template'
document.core_properties.comments = ''
document.save(out / 'Didarian_Research_Template.docx')
print('Created the DOCX source; .doc is produced separately by genuine Word format conversion.')
